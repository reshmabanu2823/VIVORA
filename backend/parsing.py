"""Read a PDF / DOCX / pasted text and split it into labelled sections."""
import io
import re

MIN_SECTION_WORDS = 40
MAX_SECTION_WORDS = 900
CHUNK_WORDS = 700
CHUNK_OVERLAP = 80

KEYWORD_HEADINGS = {
    "abstract", "introduction", "problem statement", "objectives", "scope", "project scope",
    "methodology", "proposed system", "system architecture", "architecture", "implementation",
    "results", "results and discussion", "challenges", "limitations", "conclusion",
    "future scope", "references", "literature review", "background", "acknowledgement",
}
SKIP_TITLES = {"references", "acknowledgement", "acknowledgements", "index", "table of contents",
               "contents", "bibliography", "declaration", "certificate"}

HEADING_RE = re.compile(r"^(?:(?:chapter|section)\s+)?\d+(?:\.\d+)*[.)]?\s+[A-Za-z][^\n]{2,70}$", re.I)


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        if not text.strip():
            raise ValueError("No readable text found in this PDF (it may be scanned). Paste the text instead.")
        return text
    if name.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
        if not text.strip():
            raise ValueError("This DOCX has no readable text.")
        return text
    raise ValueError("Unsupported file type. Upload a .pdf or .docx, or paste text.")


def _is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80:
        return False
    if HEADING_RE.match(s):
        return True
    if s.lower().rstrip(":") in KEYWORD_HEADINGS:
        return True
    return s.isupper() and 3 <= len(s) <= 60 and len(s.split()) <= 8


def _nice(title: str) -> str:
    """Keep the heading as the student wrote it; only fix ALL-CAPS headings."""
    title = title.strip()
    return title.title() if title.isupper() else title


def _words(text: str) -> list[str]:
    return text.split()


def _chunk(text: str, size: int = CHUNK_WORDS, overlap: int = CHUNK_OVERLAP) -> list[str]:
    words = _words(text)
    if len(words) <= size:
        return [text.strip()]
    chunks, start = [], 0
    while start < len(words):
        chunks.append(" ".join(words[start:start + size]))
        if start + size >= len(words):
            break
        start += size - overlap
    return chunks


def split_sections(text: str) -> list[dict]:
    text = re.sub(r"\r\n?", "\n", text).strip()
    if not text:
        return []

    raw: list[tuple[str, list[str]]] = []
    title, buf = "Overview", []
    found_heading = False
    for line in text.split("\n"):
        if _is_heading(line):
            found_heading = True
            if buf:
                raw.append((title, buf))
            title, buf = _nice(re.sub(r"^\d+(?:\.\d+)*[.)]?\s+", "", line.strip()).rstrip(":")), []
        else:
            buf.append(line)
    raw.append((title, buf))

    merged: list[list] = []
    for t, lines in raw:
        body = "\n".join(lines).strip()
        wc = len(_words(body))
        if wc == 0:
            continue
        if wc < MIN_SECTION_WORDS and merged:
            merged[-1][1] += "\n" + body          # fold tiny section into the previous one
        else:
            merged.append([t, body])

    if not found_heading or len(merged) <= 1:
        total = merged[0][1] if merged else text
        if len(_words(total)) > MAX_SECTION_WORDS:
            merged = [[f"Part {i + 1}", c] for i, c in enumerate(_chunk(total))]

    sections: list[dict] = []
    for t, body in merged:
        pieces = _chunk(body, MAX_SECTION_WORDS, CHUNK_OVERLAP) if len(_words(body)) > MAX_SECTION_WORDS else [body]
        for k, piece in enumerate(pieces):
            label = t if len(pieces) == 1 else f"{t} (part {k + 1})"
            sections.append({
                "id": len(sections) + 1,
                "title": label,
                "text": piece,
                "word_count": len(_words(piece)),
                "source": "REPORT"
            })
    return sections


def eligible_sections(sections: list[dict]) -> list[dict]:
    ok = [s for s in sections if s["word_count"] >= MIN_SECTION_WORDS and s["title"].lower() not in SKIP_TITLES]
    for s in (ok or sections):
        s.setdefault("source", "REPORT")
    return ok or sections


# ---------- Code parsing & Cross-verification ----------
import zipfile
from pathlib import Path

CODE_EXTENSIONS = {
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".cpp", ".c", ".h", ".hpp",
    ".cs", ".go", ".rs", ".php", ".rb", ".sql", ".html", ".css", ".json",
    ".yaml", ".yml", ".toml", ".sh", ".bash", ".kt", ".swift", ".dart", ".scala"
}

IGNORE_DIRS = {
    "node_modules", ".git", "__pycache__", ".venv", "venv", "env", ".idea",
    ".vscode", "dist", "build", "target", "vendor", ".pytest_cache", ".next",
    "coverage", ".mypy_cache"
}

VERIFIABLE_TECHNOLOGIES = [
    ("MongoDB", ["mongodb", "pymongo", "mongoose", "mongo"]),
    ("PostgreSQL", ["postgres", "postgresql", "psycopg2", "asyncpg", "pg"]),
    ("MySQL", ["mysql", "mysqlclient", "pymysql"]),
    ("SQLite", ["sqlite", "sqlite3"]),
    ("Redis", ["redis", "ioredis", "aioredis"]),
    ("Firebase", ["firebase", "firestore"]),
    ("Cassandra", ["cassandra"]),
    ("Neo4j", ["neo4j"]),
    ("Elasticsearch", ["elasticsearch", "elastic"]),
    ("JWT", ["jwt", "jsonwebtoken", "pyjwt", "bearer token", "access_token"]),
    ("OAuth", ["oauth", "oauth2", "passport", "auth0"]),
    ("bcrypt", ["bcrypt", "passlib", "hashpw"]),
    ("RAG", ["rag", "retrieval", "vector", "embed", "chroma", "pinecone", "faiss", "weaviate", "qdrant", "langchain", "llamaindex"]),
    ("Vector Database", ["chroma", "pinecone", "faiss", "weaviate", "qdrant", "pgvector"]),
    ("LangChain", ["langchain"]),
    ("LlamaIndex", ["llamaindex", "llama_index"]),
    ("OpenAI API", ["openai", "chatcompletion"]),
    ("Anthropic API", ["anthropic", "claude"]),
    ("PyTorch", ["torch", "pytorch"]),
    ("TensorFlow", ["tensorflow", "keras"]),
    ("Scikit-learn", ["sklearn", "scikit-learn"]),
    ("Docker", ["dockerfile", "docker-compose", "docker"]),
    ("Kubernetes", ["kubernetes", "k8s"]),
    ("Kafka", ["kafka", "confluent"]),
    ("RabbitMQ", ["rabbitmq", "pika", "amqp"]),
    ("Celery", ["celery"]),
    ("GraphQL", ["graphql", "apollo", "strawberry", "graphene"]),
    ("WebSockets", ["websocket", "websockets", "socket.io", "ws"]),
    ("FastAPI", ["fastapi"]),
    ("Flask", ["flask"]),
    ("Django", ["django"]),
    ("React", ["react", "useState", "useEffect"]),
    ("Vue", ["vue"]),
    ("Next.js", ["next"]),
    ("Express", ["express"]),
]


def extract_code_files(filename: str, data: bytes) -> list[dict]:
    """Extract code files from a .zip archive, individual code file, or pasted code text."""
    files = []
    name = (filename or "").lower()

    if name.endswith(".zip"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for item in z.infolist():
                    if item.is_dir():
                        continue
                    p = Path(item.filename)
                    parts = p.parts
                    if any(part in IGNORE_DIRS or part.startswith(".") for part in parts[:-1]):
                        continue
                    if p.name.startswith("."):
                        continue
                    if p.suffix.lower() not in CODE_EXTENSIONS:
                        continue
                    if item.file_size > 250 * 1024:
                        continue
                    try:
                        content_bytes = z.read(item.filename)
                        text = content_bytes.decode("utf-8", errors="replace")
                    except Exception:
                        continue
                    if text.strip():
                        files.append({
                            "path": item.filename,
                            "text": text,
                            "lines": len(text.splitlines()),
                            "size": len(content_bytes)
                        })
        except zipfile.BadZipFile:
            raise ValueError("Invalid or corrupted .zip file.")

    elif any(name.endswith(ext) for ext in CODE_EXTENSIONS) or name in ("code_upload", "pasted_code", "source_code"):
        text = data.decode("utf-8", errors="replace")
        if text.strip():
            files.append({
                "path": filename or "main.py",
                "text": text,
                "lines": len(text.splitlines()),
                "size": len(data)
            })
    else:
        try:
            text = data.decode("utf-8", errors="replace")
            if text.strip():
                files.append({
                    "path": filename or "source_code.py",
                    "text": text,
                    "lines": len(text.splitlines()),
                    "size": len(data)
                })
        except Exception:
            raise ValueError("Unsupported code file format. Please upload a .zip archive or source code file.")

    if not files:
        raise ValueError("No readable source code files found.")

    return files


def split_code_sections(code_files: list[dict], start_id: int = 1) -> list[dict]:
    """Convert code files into structured code sections."""
    sections = []
    cur_id = start_id
    for f in code_files:
        path = f["path"]
        body = f["text"].strip()
        words = body.split()
        if len(words) > MAX_SECTION_WORDS:
            pieces = _chunk(body, MAX_SECTION_WORDS, CHUNK_OVERLAP)
            for k, piece in enumerate(pieces):
                sections.append({
                    "id": cur_id,
                    "title": f"Code: {path} (part {k + 1})",
                    "file_path": path,
                    "text": piece,
                    "word_count": len(piece.split()),
                    "source": "CODE"
                })
                cur_id += 1
        else:
            sections.append({
                "id": cur_id,
                "title": f"Code: {path}",
                "file_path": path,
                "text": body,
                "word_count": len(words),
                "source": "CODE"
            })
            cur_id += 1
    return sections


def detect_cross_verification(report_sections: list[dict], code_files: list[dict]) -> dict:
    """
    Compare claims in the project report against actual source code files.
    Identifies verified features and mismatches.
    """
    full_report_text = "\n\n".join(s["text"] for s in report_sections)
    all_code_text = "\n\n".join(f["text"] for f in code_files).lower()
    all_code_paths = " ".join(f["path"].lower() for f in code_files)

    verified = []
    mismatches = []

    for tech_name, aliases in VERIFIABLE_TECHNOLOGIES:
        # Check if tech is claimed in report
        pattern = rf"\b{re.escape(tech_name)}\b"
        if not re.search(pattern, full_report_text, re.IGNORECASE):
            continue

        # Extract report excerpt sentence containing the claim
        claim_excerpt = ""
        for sec in report_sections:
            for sent in re.split(r"[.!?]\n|\.\s+", sec["text"]):
                if re.search(pattern, sent, re.IGNORECASE):
                    claim_excerpt = sent.strip()
                    break
            if claim_excerpt:
                break
        if not claim_excerpt:
            claim_excerpt = f"Report mentions {tech_name}"

        # Search code for evidence
        found_in_file = None
        for f in code_files:
            file_content_lower = f["text"].lower()
            file_path_lower = f["path"].lower()
            if any(alias in file_content_lower or alias in file_path_lower for alias in aliases):
                found_in_file = f["path"]
                break

        if found_in_file:
            verified.append({
                "tech": tech_name,
                "claim": f"Report mentions {tech_name}",
                "evidence_in_code": found_in_file,
                "status": "verified"
            })
        else:
            mismatches.append({
                "tech": tech_name,
                "claim": f"Report claims {tech_name} is used/implemented",
                "report_excerpt": claim_excerpt,
                "status": "mismatch",
                "neutral_question": f"Your report mentions {tech_name}. Can you show how {tech_name} is implemented in the project?"
            })

    return {
        "has_code": bool(code_files),
        "total_files": len(code_files),
        "verified": verified,
        "mismatches": mismatches
    }
